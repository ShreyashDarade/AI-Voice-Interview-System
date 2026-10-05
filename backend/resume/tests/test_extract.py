import io
import re

import pymupdf
import pytest
from PIL import Image, ImageDraw

from resume import EmptyResume, ResumeParseError, parse_resume
from resume import extract as X
from resume.tests.helpers import PdfBuilder, docx_bytes, simple_pdf
from resume.util import Deadline, normalize_line
from resume.validate import Limits


def extract(pdf, ocr=True):
    return X.extract_pdf(pdf, Limits(), Deadline(30), ocr=ocr)


def texts(doc):
    return [l.text for l in doc.lines]


# ------------------------------------------------------------------ normalisation
def test_normalize_line_handles_ligatures_nbsp_quotes_zero_width_and_controls():
    assert normalize_line("ﬁnal ﬂow ​zero⁠width “q” ‘s’ \x07bell") == "final flow zerowidth \"q\" 's' bell"
    assert normalize_line("Ｐｙｔｈｏｎ ① ½") == "Python 1 1⁄2"
    assert normalize_line(" item") == "• item"


# ------------------------------------------------------------------ PDF font metadata
def test_pdf_lines_carry_font_size_bold_and_position():
    b = PdfBuilder()
    b.flow(50, 60, [("BIG NAME", 24, True), ("small body text here", 10), ("Bold heading", 12, True)])
    doc = extract(b.bytes())
    by = {l.text: l for l in doc.lines}
    assert by["BIG NAME"].font_size == pytest.approx(24, abs=0.1) and by["BIG NAME"].is_bold
    assert by["small body text here"].font_size == pytest.approx(10, abs=0.1) and not by["small body text here"].is_bold
    assert by["BIG NAME"].is_upper and not by["Bold heading"].is_upper
    assert by["BIG NAME"].bbox[1] < by["small body text here"].bbox[1]
    assert by["BIG NAME"].ry0 < 0.2


def test_same_row_fragments_are_merged_with_tab_separator():
    b = PdfBuilder()
    b.text(50, 100, "Software Engineer, Acme Corp", 11, True)
    b.text(430, 100, "Jan 2020 - Present", 11)
    doc = extract(b.bytes())
    assert len(doc.lines) == 1
    assert "\t" in doc.lines[0].text and doc.lines[0].text.endswith("Jan 2020 - Present")


# ------------------------------------------------------------------ wrapping / hyphenation
def test_wrapped_bullet_lines_are_joined_and_dehyphenated():
    b = PdfBuilder()
    b.text(50, 100, "• Led the develop-", 10)
    b.text(62, 113, "ment of a platform that processes payments", 10)
    b.text(62, 126, "for millions of customers every day", 10)
    b.text(50, 150, "• Second bullet stays separate", 10)
    doc = extract(b.bytes())
    assert [l.is_bullet for l in doc.lines] == [True, True]
    assert doc.lines[0].text == "Led the development of a platform that processes payments for millions of customers every day"


def test_compound_prefix_hyphen_is_kept_when_wrapping():
    b = PdfBuilder()
    b.text(50, 100, "• Built the full-", 10)
    b.text(62, 113, "stack web application for the team", 10)
    doc = extract(b.bytes())
    assert doc.lines[0].text == "Built the full-stack web application for the team"


def test_separate_bullets_and_headings_not_merged():
    b = PdfBuilder()
    b.flow(50, 60, [("EXPERIENCE", 12, True), "• one bullet here", "• another bullet", ("EDUCATION", 12, True), "university text"])
    doc = extract(b.bytes())
    assert texts(doc) == ["EXPERIENCE", "one bullet here", "another bullet", "EDUCATION", "university text"]


# ------------------------------------------------------------------ columns
def _two_column(left_x=40, right_x=320):
    b = PdfBuilder()
    b.text(40, 55, "JOHN DOE", 22, True)
    b.text(40, 75, "john@doe.com | +1 415 555 0132", 10)
    b.flow(left_x, 120, [("LEFT ONE", 11, True)] + [f"left column line {i} with some text" for i in range(8)], width=230)
    b.flow(right_x, 120, [("RIGHT ONE", 11, True)] + [f"right column line {i} with some text" for i in range(8)], width=230)
    return b.bytes()


def test_two_column_reading_order_left_then_right_with_header_on_top():
    doc = extract(_two_column())
    t = texts(doc)
    assert t[0] == "JOHN DOE"
    assert t[1].startswith("john@doe.com")
    left = [i for i, s in enumerate(t) if s.startswith("left column") or s == "LEFT ONE"]
    right = [i for i, s in enumerate(t) if s.startswith("right column") or s == "RIGHT ONE"]
    assert max(left) < min(right)
    assert [l.column for l in doc.lines if l.text.startswith("left")] == [1] * 8
    assert [l.column for l in doc.lines if l.text.startswith("right")] == [2] * 8
    # no interleaving: left lines are in order
    assert [t[i] for i in left][1:] == [f"left column line {i} with some text" for i in range(8)]


def test_single_column_page_is_not_split():
    b = PdfBuilder()
    b.flow(50, 60, [f"A fairly long single column line number {i} that spans most of the page width easily enough" for i in range(20)])
    doc = extract(b.bytes())
    assert all(l.column == 0 for l in doc.lines)
    assert texts(doc)[3].startswith("A fairly long single column line number 3")


def test_right_aligned_dates_do_not_trigger_column_mode():
    b = PdfBuilder()
    for i in range(8):
        y = 100 + i * 40
        b.text(50, y, f"Engineer {i}, Company {i}", 11, True)
        b.text(440, y, f"Jan 20{10 + i} - Dec 20{11 + i}", 10)
        b.text(50, y + 14, f"• Delivered a long bullet describing the work in detail number {i} for the business", 10)
    doc = extract(b.bytes())
    assert all(l.column == 0 for l in doc.lines)
    assert "Engineer 3, Company 3\tJan 2013 - Dec 2014" == texts(doc)[3 * 2 - 0] or any(
        t == "Engineer 3, Company 3\tJan 2013 - Dec 2014" for t in texts(doc))


def test_sidebar_layout_reads_sidebar_then_main():
    b = PdfBuilder()
    b.rect(0, 0, 170, 842, fill=(0.2, 0.25, 0.4))
    b.flow(15, 80, [("CONTACT", 11, True, (1, 1, 1))] + [(f"sidebar item {i}", 10, False, (1, 1, 1)) for i in range(8)], width=140)
    b.flow(200, 80, [("MAIN TITLE", 14, True)] + [f"main column content line {i} goes here" for i in range(8)], width=330)
    doc = extract(b.bytes())
    t = texts(doc)
    assert t.index("sidebar item 7") < t.index("MAIN TITLE")
    assert not doc.hidden                                   # white text on dark box is legitimate


# ------------------------------------------------------------------ header/footer removal
def test_repeated_headers_footers_and_page_numbers_removed():
    b = PdfBuilder()
    for p in range(3):
        if p:
            b.new_page()
        b.text(50, 30, "Jane Doe - Resume 2024", 9)
        b.flow(50, 100, [f"Body line {p}-{i} with unique content for this page" for i in range(5)])
        b.text(50, 815, "Confidential resume of Jane", 8)
        b.text(280, 830, f"Page {p + 1} of 3", 8)
    doc = extract(b.bytes())
    t = texts(doc)
    assert not any(s.startswith("Page ") for s in t)
    assert not any("Confidential" in s for s in t)
    assert [s for s in t if s.startswith("Jane Doe - Resume")] == ["Jane Doe - Resume 2024"]     # first page keeps its header
    assert sum(1 for s in t if s.startswith("Body line")) == 15


# ------------------------------------------------------------------ links
def test_pdf_link_annotations_extracted_with_labels():
    b = PdfBuilder()
    b.text(50, 100, "My GitHub profile", 11)
    rect = pymupdf.Rect(50, 88, 160, 104)
    b.page.insert_link({"kind": pymupdf.LINK_URI, "from": rect, "uri": "https://github.com/janedoe"})
    b.page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 200, 100, 220), "uri": "mailto:jane@x.org"})
    doc = extract(b.bytes())
    assert {"url": "https://github.com/janedoe", "label": "My GitHub profile", "page": 1} in doc.links
    assert any(l["url"] == "mailto:jane@x.org" for l in doc.links)


def test_hyperlink_only_visible_through_annotation_reaches_contact_links():
    b = PdfBuilder()
    b.flow(50, 60, [("Jane Doe", 20, True), "LinkedIn | GitHub"])
    b.page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 88, 90, 104), "uri": "https://www.linkedin.com/in/jane-doe"})
    b.page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(100, 88, 140, 104), "uri": "https://github.com/jane-doe"})
    r = parse_resume(b.bytes(), "cv.pdf")
    types = {l.type for l in r.contact.links}
    assert {"linkedin", "github"} <= types


# ------------------------------------------------------------------ hidden text
def test_white_on_white_text_removed_and_reported():
    b = PdfBuilder()
    b.flow(50, 60, ["Jane Doe", "jane@x.org", ("SKILLS", 12, True), "Python, SQL"])
    b.text(50, 300, "Kubernetes Terraform Rust Haskell Scala Kafka Spark Airflow expert in everything", 10, False, (1, 1, 1))
    doc = extract(b.bytes())
    assert "Kubernetes" not in " ".join(texts(doc))
    assert doc.hidden and doc.hidden[0]["reason"] == "white_text"


def test_near_white_text_is_also_hidden():
    b = PdfBuilder()
    b.text(50, 100, "Visible resume text", 11)
    b.text(50, 200, "hidden near white keywords", 11, False, (0.97, 0.97, 0.97))
    doc = extract(b.bytes())
    assert [h["reason"] for h in doc.hidden] == ["white_text"]


def test_tiny_font_and_offpage_text_hidden():
    b = PdfBuilder()
    b.text(50, 100, "Visible resume text", 11)
    b.text(50, 300, "microscopic keyword stuffing", 1.5)
    b.text(700, 400, "offpage text far to the right", 10)
    b.text(50, -30, "above the page", 10)
    doc = extract(b.bytes())
    reasons = sorted(h["reason"] for h in doc.hidden)
    assert reasons == ["off_page", "off_page", "tiny_font"]
    assert texts(doc) == ["Visible resume text"]


def test_white_text_on_dark_background_is_not_hidden():
    b = PdfBuilder()
    b.rect(30, 80, 560, 140, fill=(0.1, 0.1, 0.3))
    b.text(50, 110, "WHITE ON NAVY BANNER", 14, True, (1, 1, 1))
    b.text(50, 200, "ordinary text", 11)
    doc = extract(b.bytes())
    assert "WHITE ON NAVY BANNER" in texts(doc) and not doc.hidden


def test_hidden_text_excluded_from_skills_and_flagged_high_when_long():
    b = PdfBuilder()
    b.flow(50, 60, [("Jane Doe", 20, True), "jane@x.org | +1 415 555 0132", ("SKILLS", 12, True), "Python, SQL"])
    stuffed = "Kubernetes Terraform Rust Haskell Scala Kafka Spark Airflow Snowflake Databricks Pulumi Ansible " * 2
    b.text(50, 400, stuffed, 10, False, (1, 1, 1))
    r = parse_resume(b.bytes(), "cv.pdf")
    skills = {s.name for s in r.skills}
    assert "Python" in skills and "Kubernetes" not in skills and "Terraform" not in skills
    flag = r.integrity.get("hidden_text")
    assert flag and flag.severity == "high" and flag.details["chars"] > 50
    assert len(flag.details["sample"]) <= 200 and "Kubernetes" in flag.details["sample"]
    assert "hidden_text_removed" in r.warnings
    assert r.integrity.risk_score >= 0.4


def test_short_hidden_text_is_lower_severity():
    b = PdfBuilder()
    b.flow(50, 60, ["Jane Doe", "jane@x.org"])
    b.text(50, 300, "Kubernetes", 10, False, (1, 1, 1))
    r = parse_resume(b.bytes(), "cv.pdf")
    assert r.integrity.get("hidden_text").severity == "low"


# ------------------------------------------------------------------ OCR path
def _image_only_pdf():
    img = Image.new("RGB", (800, 1000), "white")
    ImageDraw.Draw(img).text((50, 50), "scanned text", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=buf.getvalue())
    return doc.tobytes()


def test_scanned_pdf_without_ocr_engine_raises_clear_error(monkeypatch):
    monkeypatch.setattr(X, "_load_ocr", lambda: None)
    with pytest.raises(EmptyResume) as ei:
        parse_resume(_image_only_pdf(), "scan.pdf")
    assert "OCR" in str(ei.value)


def test_scanned_pdf_with_ocr_disabled_raises(monkeypatch):
    monkeypatch.setattr(X, "_load_ocr", lambda: (_ for _ in ()).throw(AssertionError("must not be called")))
    with pytest.raises(EmptyResume):
        parse_resume(_image_only_pdf(), "scan.pdf", ocr=False)


def test_scanned_pdf_uses_ocr_when_available(monkeypatch):
    calls = []

    def fake_ocr(png, timeout=None):
        calls.append(png[:4])
        return "Jane Doe\njane@doe.dev\n+1 415 555 0132\nEXPERIENCE\nEngineer, Acme Inc  Jan 2019 - Dec 2020\n- Built Python and Docker services\nSKILLS\nPython, Docker\n"

    monkeypatch.setattr(X, "_load_ocr", lambda: fake_ocr)
    r = parse_resume(_image_only_pdf(), "scan.pdf")
    assert calls == [b"\x89PNG"]
    assert r.source.ocr_used and r.source.extraction_method == "pymupdf-dict+ocr"
    assert r.contact.email == "jane@doe.dev" and r.contact.name == "Jane Doe"
    assert r.total_experience_months == 24
    assert "Docker" in {s.name for s in r.skills}


def test_ocr_failure_is_warning_not_crash_when_some_text_exists(monkeypatch):
    def boom(png, timeout=None):
        raise RuntimeError("tesseract exploded")

    monkeypatch.setattr(X, "_load_ocr", lambda: boom)
    b = PdfBuilder()
    b.text(50, 100, "Jane Doe jane@x.org", 12)               # < 200 chars -> sparse
    r = parse_resume(b.bytes(), "x.pdf")
    assert "ocr_failed" in r.warnings and r.contact.email == "jane@x.org"


def test_sparse_text_without_ocr_warns(monkeypatch):
    monkeypatch.setattr(X, "_load_ocr", lambda: None)
    b = PdfBuilder()
    b.text(50, 100, "Jane Doe jane@x.org", 12)
    r = parse_resume(b.bytes(), "x.pdf")
    assert "ocr_unavailable_scanned_pdf" in r.warnings


# ------------------------------------------------------------------ DOCX
def _docx_with_features():
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Header Name | header@corp.io"
    d.sections[0].footer.paragraphs[0].text = "Page 1"
    d.add_heading("Jane Q. Doe", level=0)
    p = d.add_paragraph("Profile links: ")
    # real hyperlink through relationships
    rid = d.part.relate_to("https://github.com/janedoe", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "my github"
    r.append(t)
    h.append(r)
    p._p.append(h)
    d.add_heading("EXPERIENCE", level=1)
    tbl = d.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Lead Engineer, Acme"
    tbl.cell(0, 1).text = "Jan 2020 - Present"
    tbl.cell(1, 0).text = "Layout cell A"
    tbl.cell(1, 1).text = "Layout cell B"
    d.add_paragraph("Built things in Python", style="List Bullet")
    vis = d.add_paragraph("Visible sentence. ")
    hid = vis.add_run("HIDDENVANISH keyword stuffing Kubernetes")
    hid.font.hidden = True
    white = vis.add_run(" WHITERUN secret Terraform")
    white.font.color.rgb = RGBColor(255, 255, 255)
    tiny = vis.add_run(" TINYRUN Rust")
    tiny.font.size = Pt(1)
    d.add_paragraph("Line one").add_run().add_break(WD_BREAK.LINE)
    return d


def test_docx_tables_headers_hyperlinks_bullets_and_hidden_runs():
    import docx

    d = _docx_with_features()
    buf = io.BytesIO()
    d.save(buf)
    doc = X.extract_docx(buf.getvalue(), Limits(), Deadline(30))
    t = texts(doc)
    assert t[0] == "Header Name | header@corp.io"                  # header first
    assert "Lead Engineer, Acme\tJan 2020 - Present" in t         # 2-cell row -> one line
    assert "Layout cell A\tLayout cell B" in t
    assert any(l.is_bullet and l.text == "Built things in Python" for l in doc.lines)
    assert any(l.heading_style and l.text == "EXPERIENCE" for l in doc.lines)
    assert {"url": "https://github.com/janedoe", "label": "my github", "page": 1} in doc.links
    joined = " ".join(t)
    assert "HIDDENVANISH" not in joined and "WHITERUN" not in joined and "TINYRUN" not in joined
    assert "Visible sentence." in joined
    reasons = {h["reason"] for h in doc.hidden}
    assert reasons == {"hidden_run"} and len(doc.hidden) == 3
    # name is the biggest font on the doc: title style
    assert max(doc.lines, key=lambda l: l.font_size).text == "Jane Q. Doe"


def test_docx_text_box_content_is_extracted():
    import docx
    from lxml import etree

    d = docx.Document()
    d.add_paragraph("Body paragraph")
    p = d.add_paragraph()
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><w:txbxContent>'
        '<w:p><w:r><w:t>Textbox: jane@box.io</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>Textbox: jane@box.io</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>'
    )
    p._p.append(etree.fromstring(xml))
    buf = io.BytesIO()
    d.save(buf)
    doc = X.extract_docx(buf.getvalue(), Limits(), Deadline(30))
    assert [t for t in texts(doc) if "jane@box.io" in t] == ["Textbox: jane@box.io"]      # Choice only, no duplicate from Fallback
    r = parse_resume(buf.getvalue(), "x.docx")
    assert r.contact.email == "jane@box.io"


def test_docx_white_text_on_dark_shading_is_kept():
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    d = docx.Document()
    p = d.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "1F2A44")
    ppr.append(shd)
    r = p.add_run("WHITE ON DARK HEADER")
    r.font.color.rgb = RGBColor(255, 255, 255)
    buf = io.BytesIO()
    d.save(buf)
    doc = X.extract_docx(buf.getvalue(), Limits(), Deadline(30))
    assert "WHITE ON DARK HEADER" in texts(doc) and not doc.hidden
